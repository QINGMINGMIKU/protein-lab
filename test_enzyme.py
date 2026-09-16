"""test_enzyme.py — 酶活存档契约 v2 + 复制回放（restore）回归测试（assert 脚本）

跑法（必须用 venv python）：
    .venv/Scripts/python.exe test_enzyme.py

覆盖：
1. 纯函数 enzyme_time_grid / time_axis_to_indices —— 网格重建与秒值对→下标（含最近邻/clamp/越界）
2. 存档契约 v2（POST /api/experiments/from-calculation）—— params.wells 无逐点数据但留 od_range/mw；
   raw payload 带 analysis_version / calc_type / params / source_file / 全量 wells
3. POST /api/enzyme/restore —— 新形状回填 / 旧形状（v1 无 params 槽）兼容合成 / 空快照 400 /
   版本不符带 version_warning 不阻断
4. 不可变性：resave 重挂后旧 raw 逐字不变、新 raw 为新契约
5. 归档导出「作图数据」Sheet 从 raw 读全量（raw 缺失回退 params）

数据安全：数据库用临时目录，不触碰生产库（见 CLAUDE.md 测试规范）。
"""
import os, sys, importlib, tempfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from calculators import (ENZYME_ANALYSIS_VERSION, enzyme_time_grid,
                         time_axis_to_indices)

TMP = tempfile.mkdtemp(prefix="protein_lab_test_")

# ── 1. 纯函数：时间网格 + 秒值对→下标 ──
grid = enzyme_time_grid({"temps": [20.0, 0.0, 10.0, 10.0]}, {})
assert grid == [0.0, 10.0, 20.0], f"网格应去重升序: {grid}"
grid2 = enzyme_time_grid({}, {"A1": {"times": [0, 5, 10]}})   # meta 无 temps → 回落首孔
assert grid2 == [0.0, 5.0, 10.0], f"缺 temps 应回落首孔 times: {grid2}"
assert enzyme_time_grid({}, {}) == [], "无网格应返回空表"
assert enzyme_time_grid({"temps": ["x", 5, None]}, {}) == [5.0], "非数值网格点应跳过"

G = [0.0, 10.0, 20.0, 30.0]
assert time_axis_to_indices(G, [10.0, 30.0]) == (1, 3), "落点精确命中"
assert time_axis_to_indices(G, [None, None]) == (0, 3), "缺省回退整区间"
assert time_axis_to_indices(G, None) == (0, 3), "None 回退整区间"
assert time_axis_to_indices(G, [30.0, 10.0]) == (1, 3), "逆序入参应交换"
assert time_axis_to_indices(G, [11.0, 12.0]) == (1, 1), "最近邻：11→10(1)，12→10(1)"
assert time_axis_to_indices(G, [-99.0, 999.0]) == (0, 3), "越界 clamp 到两端"
assert time_axis_to_indices(G, ["a", "b"]) == (0, 3), "非数值回退整区间"
assert time_axis_to_indices([], [1, 2]) == (0, 0), "空网格不崩"
print("1. enzyme_time_grid / time_axis_to_indices OK")

# ── 2. 隔离临时库（顺序不能乱，见 CLAUDE.md 测试规范）──
import models
importlib.reload(models)  #  会把 DB_PATH 重置为真实路径
models.DB_PATH = os.path.join(TMP, "test.db")
models.init_db()
from app import app
client = app.test_client()

TT = [0.0, 5.0, 10.0, 15.0, 20.0]
OD_A1 = [0.1000, 0.2000, 0.3000, 0.4000, 0.5000]
OD_A2 = [0.0500, 0.0500, 0.0500, 0.0500, 0.0500]
# 手工处理参数（= 前端 enzymeParams() 的产物）
EPARAMS = {
    "calc_type": "enzyme",
    "meta": {"sample": "S1", "wavelength": 340, "temps": TT},
    "source_file": "20260916_120000.xlsx",
    "time_axis": [5.0, 15.0],
    "sub_blank": True, "show_blank": False, "align_start": True,
    "align_end": False, "error_bar": "sem", "group": True,
}
PWELLS = {
    "A1": {"name": "WT", "ref": "", "group": "WT", "protein_id": None, "mw": 27000,
           "conc_ng_ml": 100.0, "conc_uM": 3.7, "fit": {"slope": 0.02, "r2": 0.99},
           "od_range": ["0.2000", "0.4000"]},
    "A2": {"name": "Blank", "ref": "blank", "group": "", "protein_id": None, "mw": None,
           "conc_ng_ml": None, "conc_uM": None, "fit": None, "od_range": ["0.0500", "0.0500"]},
}
RAW_WELLS = {
    "A1": {"name": "WT", "ref": "", "times": TT, "od": OD_A1},
    "A2": {"name": "Blank", "ref": "blank", "times": TT, "od": OD_A2},
}

r = client.post("/api/experiments/from-calculation", json={
    "title": "", "exp_type": "酶活测定", "date": "2026-09-16", "calc_type": "enzyme",
    "calc_params": {**EPARAMS, "wells": PWELLS, "well_count": 2},
    "calc_result": {},
    "raw_snapshots": [{
        "data_type": "enzyme_traces",
        "payload": {"analysis_version": ENZYME_ANALYSIS_VERSION, "calc_type": "enzyme",
                    "params": EPARAMS, "source_file": EPARAMS["source_file"],
                    "meta": EPARAMS["meta"], "wells": RAW_WELLS},
    }],
})
assert r.status_code == 201, f"存档应 201: {r.status_code} {r.get_json()}"
eid = r.get_json()["id"]

# 2a. params.wells 无逐点数据，但保留 od_range / mw（契约规则 A）
e = models.exp_get(eid)
pw = e["params"]["wells"]["A1"]
assert "times" not in pw and "od" not in pw, f"params.wells 不应内嵌逐点数据: {sorted(pw)}"
assert pw["od_range"] == ["0.2000", "0.4000"], "params.wells 应保留 od_range"
assert pw["mw"] == 27000, "params.wells 应保留 mw"
assert e["params"]["sub_blank"] is True and e["params"]["error_bar"] == "sem", "手动开关应落 params"
# 2b. raw payload 契约槽齐全 + 全量数据
raws = models.exp_raw_list(eid)
assert len(raws) == 1 and raws[0]["data_type"] == "enzyme_traces", f"raw 应落 1 条: {raws}"
payload = models.exp_raw_get(raws[0]["id"])["payload"]
for k in ("analysis_version", "calc_type", "params", "source_file", "wells"):
    assert k in payload, f"payload 应带契约槽 {k}: {sorted(payload)}"
assert payload["wells"]["A1"]["times"] == TT, "raw 应存全量时间序列"
assert payload["params"]["time_axis"] == [5.0, 15.0], "payload.params 应是手动参数冻结副本"
assert payload["source_file"] == "20260916_120000.xlsx", "payload 应带源文件名"
# 2c. _raws 带 data_type（供前端按类型取最新快照）
eg = client.get(f"/api/experiments/{eid}").get_json()
assert [x["data_type"] for x in eg["_raws"]] == ["enzyme_traces"], f"_raws 应带类型: {eg['_raws']}"
assert eg["_raw_ids"] == [raws[0]["id"]], "_raw_ids 应保留（兼容旧调用方）"
print("2. 存档契约 v2（params 无逐点 / raw 全量 / _raws 带类型）OK")

# ── 3. /api/enzyme/restore ──
raw = client.get(f"/api/experiments/{eid}/raw/{raws[0]['id']}").get_json()
d = client.post("/api/enzyme/restore", json={"payload": raw["payload"]})
assert d.status_code == 200, f"restore 应 200: {d.status_code} {d.get_json()}"
d = d.get_json()
assert d["wells"]["A1"]["times"] == TT, "restore 应回全量曲线"
assert d["n_points"] == 5 and (d["time_lo"], d["time_hi"]) == (1, 3), \
    f"time_axis [5,15] → 下标 (1,3)，实为 {(d['time_lo'], d['time_hi'])}"
assert d["params"]["sub_blank"] is True and d["params"]["error_bar"] == "sem", "restore 应回手动参数"
assert "version_warning" not in d, "版本一致不应有警告"

# 3a. 旧形状 payload（v1：无 params 槽、time_axis 内嵌）→ 合成最小 params + 版本警告
legacy = {"analysis_version": "enzyme-1.0", "meta": {"temps": TT},
          "wells": RAW_WELLS, "time_axis": [5.0, 15.0]}
d2 = client.post("/api/enzyme/restore", json={"payload": legacy}).get_json()
assert d2["params"]["calc_type"] == "enzyme", "旧快照应合成 calc_type"
assert d2["params"]["time_axis"] == [5.0, 15.0], "旧快照 time_axis 应保留"
assert (d2["time_lo"], d2["time_hi"]) == (1, 3), "旧快照也应换算时间窗"
assert "version_warning" in d2, "版本不符应给 version_warning（不阻断）"
assert d2["wells"]["A1"]["od"] == OD_A1, "旧快照曲线应原样透传"
# 3b. 旧快照无 time_axis → 全区间
d2b = client.post("/api/enzyme/restore",
                  json={"payload": {**legacy, "time_axis": None}}).get_json()
assert (d2b["time_lo"], d2b["time_hi"]) == (0, 4), "无 time_axis 应回全区间"
# 3c. 空 wells / 非对象 payload → 400
r400 = client.post("/api/enzyme/restore", json={"payload": {"wells": {}}})
assert r400.status_code == 400, f"空快照应 400: {r400.status_code}"
assert "孔" in r400.get_json()["error"], f"错误文案应说明缺孔位: {r400.get_json()}"
assert client.post("/api/enzyme/restore", json={"payload": []}).status_code == 400
assert client.post("/api/enzyme/restore", json={}).status_code == 400
print("3. /api/enzyme/restore（新形状 / v1 兼容 / 400 / 版本警告）OK")

# ─ 4. 不可变性：重挂追加 raw，旧 raw 逐字不变 ──
import services
services.resave_experiment(eid, params={**EPARAMS, "wells": PWELLS}, results={},
                           raw_snapshots=[("enzyme_traces", {
                               "analysis_version": ENZYME_ANALYSIS_VERSION,
                               "calc_type": "enzyme", "params": EPARAMS,
                               "source_file": EPARAMS["source_file"],
                               "meta": EPARAMS["meta"], "wells": RAW_WELLS})])
raws2 = models.exp_raw_list(eid)
assert len(raws2) == 2, f"重挂应追加 raw，total=2: {len(raws2)}"
assert models.exp_raw_get(raws[0]["id"])["payload"] == payload, "旧 raw 不得被覆盖"
assert models.exp_raw_get(raws2[1]["id"])["payload"]["params"]["calc_type"] == "enzyme", \
    "新 raw 应为新契约形状"
# 前端按 data_type 取最新 → 拿到的是第 2 条
eg2 = client.get(f"/api/experiments/{eid}").get_json()
assert eg2["_raws"][-1]["id"] == raws2[1]["id"], "最新一条应是重挂的那条"
print("4. raw 只插不更（重挂追加 / 旧 raw 不变）OK")

# ── 5. 归档导出「作图数据」Sheet：raw 优先，缺 raw 回退 params ──
# A：有 raw（全量 5 点）；B：无 raw，params.wells 内嵌 3 点（历史形状）
EXP_PARAMS_B = {
    "calc_type": "enzyme", "meta": {"sample": "S2"},
    "wells": {"B1": {"name": "OLD", "ref": "", "fit": None,
                     "times": [5.0, 10.0, 15.0], "od": [0.2, 0.3, 0.4]}},
    "time_axis": [5.0, 15.0],
}
eB = services.create_experiment(title="无raw酶活", exp_type="酶活测定", date="2026-09-16",
                                params=EXP_PARAMS_B, results={})
from openpyxl import load_workbook
import io
r = client.get("/api/experiments/export?calc_type=enzyme")
assert r.status_code == 200, f"导出应 200: {r.status_code}"
wb = load_workbook(io.BytesIO(r.data))
assert "作图数据" in wb.sheetnames, f"应有作图数据 Sheet: {wb.sheetnames}"
ws = wb["作图数据"]
hdr = [c.value or "" for c in ws[1]]
i_a1 = next(i for i, h in enumerate(hdr) if h.endswith("WT 时间 (min)")) + 1   # 前缀=实验标题
col_a1 = [ws.cell(row=r_, column=i_a1).value for r_ in range(2, ws.max_row + 1)]
col_a1 = [v for v in col_a1 if v is not None]
assert len(col_a1) == 5, f"有 raw 应导出全量 5 点，实为 {len(col_a1)}（回归：曾被截断为区间内点数）"
i_b1 = next(i for i, h in enumerate(hdr) if h.endswith("OLD 时间 (min)")) + 1
col_b1 = [ws.cell(row=r_, column=i_b1).value for r_ in range(2, ws.max_row + 1)]
col_b1 = [v for v in col_b1 if v is not None]
assert len(col_b1) == 3, f"无 raw 应回退 params 的 3 点（不降级），实为 {len(col_b1)}"
# 单位换算：秒 → 分钟（导出列按 min 写）
assert abs(col_a1[0]) < 1e-9 and abs(col_a1[1] - 5 / 60) < 1e-3, f"时间列应换算为分钟: {col_a1[:2]}"
print(f"5. 归档导出作图数据（raw 全量 {len(col_a1)} 点 / 无 raw 回退 {len(col_b1)} 点）OK")

# ── 6. tools/strip_params_pointdata.py：历史 params 去逐点（**fail-closed**）──
# 这是一次性迁移，写的是生产库，所以它的「不删唯一副本」防线必须有回归。
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "tools"))
import strip_params_pointdata as strip

STRIP_TT = [0.0, 10.0, 20.0, 30.0]


def _mk(title, wells, raw=None, calc_type="enzyme"):
    """建一条实验；raw=None 表示不落原始快照。"""
    return services.create_experiment(
        title=title, exp_type="酶活测定", date="2026-09-16",
        params={"calc_type": calc_type, "meta": {"sample": title, "temps": STRIP_TT},
                "wells": wells},
        results={},
        raw_snapshots=[("enzyme_traces", raw)] if raw else None,
    )["id"]


_FULL_RAW = {"analysis_version": ENZYME_ANALYSIS_VERSION, "calc_type": "enzyme",
             "params": {"time_axis": [10.0, 20.0]},
             "meta": {"sample": "x", "temps": STRIP_TT},
             "wells": {"A1": {"times": STRIP_TT, "od": [0.1, 0.2, 0.3, 0.4]}}}
# 6a. 有完整 raw 覆盖 → 去逐点，且**保留** fit / od_range 等元数据
_w_pd = {"A1": {"name": "WT", "ref": "", "group": "G", "mw": 27000,
                "conc_ng_ml": 100.0, "conc_uM": 3.7, "fit": {"slope": 0.02, "r2": 0.99},
                "od_range": ["0.1000", "0.4000"],
                "times": [10.0, 20.0], "od": [0.2, 0.3]}}      # 被时间窗截断的 2 点
_e6a = _mk("去逐点A", _w_pd, _FULL_RAW)
_p6a = strip.plan_experiment(_e6a)
assert _p6a["action"] == "strip", f"有完整 raw 覆盖应 strip: {_p6a}"
models.exp_update(_e6a, params=_p6a["params"])
_w6a = models.exp_get(_e6a)["params"]["wells"]["A1"]
assert "times" not in _w6a and "od" not in _w6a, f"逐点应已删除: {sorted(_w6a)}"
for _k, _v in (("name", "WT"), ("group", "G"), ("mw", 27000), ("conc_uM", 3.7),
               ("od_range", ["0.1000", "0.4000"])):
    assert _w6a[_k] == _v, f"白名单重建漏了 {_k}: {_w6a.get(_k)!r}"
assert _w6a["fit"] == {"slope": 0.02, "r2": 0.99}, "fit 必须保留（详情页/对比都读它）"
# time_axis 取 raw 的权威值；errors 标签 time_unit 被摘掉
assert models.exp_get(_e6a)["params"]["time_axis"] == [10.0, 20.0], "time_axis 应用 raw 的值"
# 6b. 幂等：再跑一次判 skip，不重复动作
assert strip.plan_experiment(_e6a)["action"] == "skip", "去逐点后应幂等 skip"
# 6c. **无 raw** → fail（逐点是唯一副本，一个点都不能删）
_e6c = _mk("无raw", {"A1": {"name": "X", "times": STRIP_TT, "od": [0.1, 0.2, 0.3, 0.4]}})
_p6c = strip.plan_experiment(_e6c)
assert _p6c["action"] == "fail" and "唯一副本" in _p6c["reason"], f"无 raw 应 fail: {_p6c}"
# 6d. raw 点数 < params 点数 → fail（拒绝用截断数据覆盖全量）
_short_raw = {**_FULL_RAW, "wells": {"A1": {"times": [0.0, 10.0], "od": [0.1, 0.2]}}}
_e6d = _mk("raw更短", {"A1": {"name": "X", "times": STRIP_TT, "od": [0.1, 0.2, 0.3, 0.4]}}, _short_raw)
_p6d = strip.plan_experiment(_e6d)
assert _p6d["action"] == "fail" and "截断" in _p6d["reason"], f"raw 更短应 fail: {_p6d}"
# 6e. 某个孔在 raw 里不存在 → 整条 fail（逐实验原子，不部分删）
_miss_raw = {**_FULL_RAW, "wells": {"A1": _FULL_RAW["wells"]["A1"]}}
_e6e = _mk("缺孔", {"A1": {"name": "X", "times": STRIP_TT, "od": [0.1, 0.2, 0.3, 0.4]},
                    "B2": {"name": "Y", "times": STRIP_TT, "od": [0.1, 0.2, 0.3, 0.4]}}, _miss_raw)
_p6e = strip.plan_experiment(_e6e)
assert _p6e["action"] == "fail" and "B2" in _p6e["reason"], f"缺孔应 fail: {_p6e}"
# 6f. #47 场景：times 误为分钟 + meta.time_unit="min" → 一并清掉，fit 保留
_e6f = services.create_experiment(
    title="分钟脏数据", exp_type="酶活测定", date="2026-09-16",
    params={"calc_type": "enzyme", "meta": {"sample": "min", "time_unit": "min"},
            "time_axis": [0.167, 0.333],
            "wells": {"A1": {"name": "M", "fit": {"slope": 0.5}, "times": [0.0, 0.167],
                             "od": [0.1, 0.2]}}},
    results={}, raw_snapshots=[("enzyme_traces", _FULL_RAW)])["id"]
_p6f = strip.plan_experiment(_e6f)
assert _p6f["action"] == "strip", _p6f
models.exp_update(_e6f, params=_p6f["params"])
_p6f_after = models.exp_get(_e6f)["params"]
assert "time_unit" not in _p6f_after["meta"], "错误的分钟标签应被摘掉"
assert _p6f_after["time_axis"] == [10.0, 20.0], "time_axis 应换成 raw 的秒值"
assert _p6f_after["wells"]["A1"]["fit"] == {"slope": 0.5}, "存档 fit 应保留（它本来就是对的）"
# 6g. 非酶活实验不被本迁移波及
_e6g = services.create_experiment(title="BLI 不碰", exp_type="BLI",
                                  params={"calc_type": "bli_fit", "wells": {"A1": {"times": [1], "od": [1]}}},
                                  results={})["id"]
assert strip.plan_experiment(_e6g)["action"] == "skip", "非酶活实验应 skip"
assert "times" in models.exp_get(_e6g)["params"]["wells"]["A1"], "非酶活实验的 params 不得被改"
# 6h. 扫库目标只含酶活实验
assert _e6g not in strip.targets_from_db(), "targets_from_db 不应含非酶活实验"
assert _e6c in strip.targets_from_db(), "targets_from_db 应含酶活实验"
print("6. tools/strip_params_pointdata（去逐点 / 保留元数据 / 幂等 / 无raw缺孔截断三防线 / 分钟脏数据 / 不碰非酶活）OK")

import shutil
shutil.rmtree(TMP, ignore_errors=True)
print("\nALL PASSED")