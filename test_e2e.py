"""test_e2e.py — 真实样例端到端回归（assert 脚本）

跑法（必须用 venv python）：
    .venv/Scripts/python.exe test_e2e.py

覆盖三个分析模块的完整链路：
    导入 → 分析/绘图 → 保存为实验 → 读原始快照 → 重新载入（restore）→ 再分析 → 导出

样例来源：
- AKTA：fixtures/*.zip（git 跟踪，CI 可用）
- BLI：REF/20260708.csv（真实 ForteBio 导出；REF/ 被 gitignore，缺失则跳过）
- 酶活：REF/20260809_174326.xlsx（真实 TECAN 导出；缺失则跳过）

数据安全：临时库，不触碰生产 protein_lab.db。
"""
import os
import sys
import io
import importlib
import tempfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = os.path.dirname(os.path.abspath(__file__))
TMP = tempfile.mkdtemp(prefix="protein_lab_e2e_")

import models
importlib.reload(models)
models.DB_PATH = os.path.join(TMP, "test.db")
models.init_db()

from app import app
from calculators import ENZYME_ANALYSIS_VERSION

client = app.test_client()
PASSED = []
SKIPPED = []


def _latest_raw(exp_id):
    detail = client.get(f"/api/experiments/{exp_id}").get_json()
    assert detail["_raw_ids"], f"实验 {exp_id} 没有原始快照"
    rid = detail["_raw_ids"][-1]
    resp = client.get(f"/api/experiments/{exp_id}/raw/{rid}")
    assert resp.status_code == 200
    return resp.get_json()["payload"]


def _xlsx_bytes(resp):
    assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
    assert resp.data[:2] == b"PK", "导出应为 xlsx(zip) 容器"
    return resp.data


# ═══════════════════════════════════════════════════════
#  AKTA：fixtures 真实 zip（CI 也跑）
# ═══════════════════════════════════════════════════════
ZIP = os.path.join(ROOT, "fixtures", "1YPI_32 002.zip")
assert os.path.exists(ZIP), f"缺少 AKTA 样例: {ZIP}"

with open(ZIP, "rb") as fh:
    resp = client.post("/api/akta/analyze",
                       data={"file": (fh, os.path.basename(ZIP))},
                       content_type="multipart/form-data")
assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
run = resp.get_json()["runs"][0]
assert "error" not in run, run
sid = run["session_id"]
channel = (run["uv_channels"] or [run["channels"][0]["name"]])[0]
assert run["channels"], "应解析出通道"

resp = client.post("/api/akta/plot", json={
    "session_id": sid, "channel": channel, "min_height": 5, "smooth_window": 11})
assert resp.status_code == 200, resp.get_json()
plot = resp.get_json()
assert plot["image"].startswith("data:image/png;base64,")
assert isinstance(plot.get("peaks"), list) and plot["peaks"], "真实样例应检出峰"

xlsx = _xlsx_bytes(client.post("/api/akta/export", json={
    "session_id": sid, "channel": channel, "min_height": 5, "smooth_window": 11}))

resp = client.post("/api/akta/save", json={
    "session_id": sid, "title": "E2E AKTA 真实样例", "channel": channel,
    "min_height": 5, "smooth_window": 11, "source": os.path.basename(ZIP)})
assert resp.status_code == 201, resp.get_data(as_text=True)[:300]
akta_exp = resp.get_json()
assert akta_exp["results"]["n_peaks"] >= 1

payload = _latest_raw(akta_exp["id"])
assert payload["channel"]["vols"], "raw 应保存全量通道曲线"
assert payload["params"]["calc_type"] == "akta"

resp = client.post("/api/akta/restore", json={
    "payload": payload, "name": os.path.basename(ZIP)})
assert resp.status_code == 200, resp.get_json()
restored = resp.get_json()["runs"][0]
assert restored["channels"][0]["n_points"] == len(payload["channel"]["vols"]), \
    "重新载入应还原全量曲线"
resp = client.post("/api/akta/plot", json={
    "session_id": restored["session_id"], "channel": payload["channel"]["name"],
    "min_height": 5, "smooth_window": 11})
assert resp.status_code == 200 and resp.get_json()["image"].startswith("data:image/png"), \
    "重新载入后应能再次出图"
PASSED.append("AKTA")
print("1. AKTA 导入→分析→保存→重载→导出 OK")


# ═══════════════════════════════════════════════════════
#  BLI：真实 ForteBio CSV（REF 缺失则跳过）
# ═══════════════════════════════════════════════════════
BLI_CSV = os.path.join(ROOT, "REF", "20260708.csv")
if not os.path.exists(BLI_CSV):
    SKIPPED.append("BLI（缺少 REF/20260708.csv）")
    print("2. BLI 跳过（缺少 REF/20260708.csv）")
else:
    with open(BLI_CSV, "rb") as fh:
        resp = client.post("/api/bli/analyze",
                           data={"file": (fh, os.path.basename(BLI_CSV))},
                           content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
    ana = resp.get_json()
    sid = ana["session_id"]
    assert ana["samples"], "真实样例应解析出样品"
    sen = ana["samples"][0]
    sample = sen["sample"]
    assert sen["n_curves"] >= 2, "真实样例每个样品应有多条浓度曲线"

    resp = client.post("/api/bli/plot", json={
        "session_id": sid, "smooth_window": 31, "fit": True})
    assert resp.status_code == 200 and resp.get_json()["image"].startswith("data:image/png"), \
        resp.get_json()
    resp = client.post("/api/bli/plot", json={"session_id": sid, "separate": True})
    assert resp.status_code == 200 and resp.get_json()["images"], "分样本出图应返回多张"

    # 拟合：不强制 KD 数值（真实数据可能因相界/浓度过滤给 error），但端点必须正常返回
    resp = client.post("/api/bli/fit", json={"session_id": sid, "sample": sample})
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["sample"] == sample

    resp = client.post("/api/bli/save", json={
        "session_id": sid, "title": "E2E BLI 真实样例", "smooth_window": 31,
        "source": os.path.basename(BLI_CSV)})
    assert resp.status_code == 201, resp.get_data(as_text=True)[:300]
    bli_exp = resp.get_json()
    assert bli_exp["results"].get("BLI_ANALYSIS_VERSION")

    payload = _latest_raw(bli_exp["id"])
    assert payload["curves"], "raw 应保存全量曲线"
    assert payload["params"]["calc_type"] == "bli_fit"

    resp = client.post("/api/bli/restore", json={"payload": payload})
    assert resp.status_code == 200, resp.get_json()
    restored = resp.get_json()
    assert restored["n_sensors"] == len(payload["curves"]), "重新载入应还原全部曲线"
    resp = client.post("/api/bli/plot", json={
        "session_id": restored["session_id"], "smooth_window": 31})
    assert resp.status_code == 200 and resp.get_json()["image"].startswith("data:image/png"), \
        "重新载入后应能再次出图"
    _xlsx_bytes(client.post("/api/bli/export", json={"session_id": restored["session_id"]}))
    PASSED.append("BLI")
    print("2. BLI 导入→分析→保存→重载→导出 OK")


# ═══════════════════════════════════════════════════════
#  酶活：真实 TECAN xlsx（REF 缺失则跳过）
# ═══════════════════════════════════════════════════════
ENZ_XLSX = os.path.join(ROOT, "REF", "20260809_174326.xlsx")
if not os.path.exists(ENZ_XLSX):
    SKIPPED.append("酶活（缺少 REF/20260809_174326.xlsx）")
    print("3. 酶活跳过（缺少 REF/20260809_174326.xlsx）")
else:
    with open(ENZ_XLSX, "rb") as fh:
        resp = client.post("/api/enzyme/parse",
                           data={"file": (fh, os.path.basename(ENZ_XLSX))},
                           content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
    edata = resp.get_json()
    wells = edata["wells"]
    assert wells and edata["meta"]["temps"], "真实样例应有孔位与时间轴"

    resp = client.post("/api/enzyme/fit", json={"wells": wells})
    assert resp.status_code == 200, resp.get_json()
    fits = resp.get_json()["wells"]
    assert set(fits) == set(wells)

    resp = client.post("/api/enzyme/plot", json={"wells": wells, "type": "kinetics"})
    assert resp.status_code == 200 and resp.get_json()["image"].startswith("data:image/png"), \
        resp.get_json()

    # 存档参数按前端 enzymeParams() 的形状：params 只留元数据 + fit，逐点在 raw
    params_wells = {}
    for wid, w in wells.items():
        ods = [v for v in (w.get("od") or []) if v is not None]
        params_wells[wid] = {
            "name": wid, "ref": "", "mw": 0, "group": "",
            "fit": fits.get(wid),
            "od_range": [f"{min(ods):.4f}", f"{max(ods):.4f}"] if ods else [],
        }
    calc_params = {
        "meta": edata["meta"],
        "source_file": os.path.basename(ENZ_XLSX),
        "time_axis": [edata["meta"]["temps"][0], edata["meta"]["temps"][-1]],
        "wells": params_wells,
    }
    raw_payload = {
        "analysis_version": ENZYME_ANALYSIS_VERSION,
        "calc_type": "enzyme",
        "params": {"calc_type": "enzyme", **calc_params},
        "source_file": os.path.basename(ENZ_XLSX),
        "meta": edata["meta"],
        "wells": wells,
    }
    resp = client.post("/api/experiments/from-calculation", json={
        "title": "E2E 酶活真实样例", "exp_type": "酶活测定", "calc_type": "enzyme",
        "calc_params": calc_params, "calc_result": {"wells": fits},
        "raw_snapshots": [{"data_type": "enzyme_traces", "payload": raw_payload}],
    })
    assert resp.status_code == 201, resp.get_data(as_text=True)[:300]
    enz_exp = resp.get_json()
    assert enz_exp["params"]["calc_type"] == "enzyme"
    assert "times" not in next(iter(enz_exp["params"]["wells"].values())), \
        "契约规则 A：params 不应内嵌逐点数据"

    payload = _latest_raw(enz_exp["id"])
    first_well = payload["wells"][next(iter(payload["wells"]))]
    assert len(first_well["times"]) == len(edata["wells"][next(iter(edata["wells"]))]["times"]), \
        "raw 应保存全量时间序列"

    resp = client.post("/api/enzyme/restore", json={"payload": payload})
    assert resp.status_code == 200, resp.get_json()
    restored = resp.get_json()
    assert restored["n_points"] == len(edata["meta"]["temps"]), "重新载入应还原全量时间点"
    restored_wells = restored["wells"]

    # 重新载入后再拟合 + 出图
    resp = client.post("/api/enzyme/fit", json={"wells": restored_wells})
    assert resp.status_code == 200
    resp = client.post("/api/enzyme/plot", json={"wells": restored_wells, "type": "kinetics"})
    assert resp.status_code == 200 and resp.get_json()["image"].startswith("data:image/png")

    # 导出：模块导出 + 归档导出（作图数据 Sheet 应来自 raw 全量）
    _xlsx_bytes(client.post("/api/enzyme/export", json={
        "wells": restored_wells, "wells_fit": fits}))
    xlsx = _xlsx_bytes(client.get("/api/experiments/export?calc_type=enzyme"))
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx))
    assert "作图数据" in wb.sheetnames, wb.sheetnames
    ws = wb["作图数据"]
    time_col = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    time_col = [v for v in time_col if v is not None]
    assert len(time_col) == len(edata["meta"]["temps"]), \
        f"归档导出应含全量 {len(edata['meta']['temps'])} 个时间点，实为 {len(time_col)}"
    PASSED.append("酶活")
    print("3. 酶活 导入→分析→保存→重载→导出 OK")


# ═══════════════════════════════════════════════════════
#  跨模块：档案列表 → 可载入判定 → 批量导出
# ═══════════════════════════════════════════════════════
resp = client.get("/api/experiments")
assert resp.status_code == 200
rows = resp.get_json()
e2e_rows = [r for r in rows if r["title"].startswith("E2E ")]
assert len(e2e_rows) == len(PASSED), f"档案页应列出全部 E2E 实验: {[r['title'] for r in rows]}"
assert all(r["loadable"] for r in e2e_rows), \
    f"分析型实验应可载入计算工具: {[(r['title'], r['loadable']) for r in e2e_rows]}"
_xlsx_bytes(client.get("/api/experiments/export"))
print(f"4. 档案列表/可载入判定/汇总导出 OK（{len(e2e_rows)} 条）")

print(f"\n跑通模块: {', '.join(PASSED)}")
if SKIPPED:
    print("跳过（缺少本地真实样例）: " + "；".join(SKIPPED))
print("\nALL PASSED")
